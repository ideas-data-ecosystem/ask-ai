"""Ingestion: registering documents, extraction (text layer, OCR, DOCX, MD/TXT), the job queue and job processing.

The worker (app.worker) claims jobs; the API and the seed CLI only register documents and enqueue jobs.
Stored paths are relative to STORAGE_DIR so the host CLI and the containers agree on them.
"""

import hashlib
import logging
import time
import zipfile
from dataclasses import dataclass
from pathlib import Path
from uuid import UUID, uuid4

import pypdfium2 as pdfium
from pgvector import Vector
from psycopg import Connection
from psycopg.types.json import Jsonb
from psycopg_pool import ConnectionPool

from app import embedding
from app.chunking import Chunk, chunk_pages, clean_page, normalize
from app.config import Settings
from app.db import kb_scope

log = logging.getLogger("app.ingest")

ALLOWED_TYPES = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "md": "text/markdown",
    "txt": "text/plain",
}
DOC_COLUMNS = (
    "id, kb_id, filename, category, mime, size_bytes, sha256, status, page_count, ocr_pages, chunk_count, "
    "error, indexed_at, created_at"
)
OCR_MIN_CHARS = 50  # a cleaned page shorter than this is OCRed
OCR_DPI = 250
OCR_PAGE_TIMEOUT = 120  # seconds per page
OCR_MAX_PIXELS = 4000  # cap on the longer side so a huge page cannot exhaust memory
MAX_ATTEMPTS = 3
DOCX_MAX_UNCOMPRESSED = 100 << 20  # python-docx reads every part of the zip into memory
DOCX_MAX_RATIO = 100  # uncompressed / compressed, only checked above DOCX_RATIO_FLOOR (a tiny file cannot hurt)
DOCX_RATIO_FLOOR = 10 << 20


class PermanentError(Exception):
    """The document itself cannot be processed; retrying will not help."""

    permanent = True


@dataclass(frozen=True)
class Extracted:
    pages: list[str]  # cleaned text, one entry per page (a single entry when not paged)
    paged: bool
    ocr_pages: int = 0


# --- registering documents -------------------------------------------------------------------------------


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        while block := f.read(1 << 20):
            h.update(block)
    return h.hexdigest()


def resolve_stored(settings: Settings, storage_path: str) -> Path:
    """Absolute path of a stored file; refuses anything that resolves outside STORAGE_DIR."""
    root = settings.storage_dir.resolve()
    path = (root / storage_path).resolve()
    if not path.is_relative_to(root):
        raise PermanentError("Stored path escapes the storage directory")
    return path


def add_document(
    conn: Connection,
    *,
    kb_id: UUID,
    doc_id: UUID | None = None,
    filename: str,
    category: str | None,
    storage_path: str,
    mime: str | None,
    size_bytes: int,
    sha256: str,
) -> dict | None:
    """Insert a documents row and its first job in one transaction. None when the KB already has this content."""
    with conn.transaction():
        row = conn.execute(
            f"INSERT INTO documents (id, kb_id, filename, category, storage_path, mime, size_bytes, sha256) "
            f"VALUES (%s, %s, %s, %s, %s, %s, %s, %s) ON CONFLICT (kb_id, sha256) DO NOTHING RETURNING {DOC_COLUMNS}",
            [doc_id or uuid4(), kb_id, filename, category, storage_path, mime, size_bytes, sha256],
        ).fetchone()
        if row:
            conn.execute("INSERT INTO ingestion_jobs (kb_id, document_id) VALUES (%s, %s)", [kb_id, row["id"]])
    return row


# --- extraction ------------------------------------------------------------------------------------------


def extract_pdf(path: Path, max_pages: int) -> list[str]:
    """Raw text layer of every page (pdfium). A PDF over `max_pages` fails before any page is read."""
    try:
        pdf = pdfium.PdfDocument(str(path))
    except pdfium.PdfiumError as e:
        log.warning("cannot open %s: %s", path.name, e)
        raise PermanentError("Cannot open PDF (not a valid or readable PDF file)") from e
    try:
        if len(pdf) > max_pages:
            raise PermanentError(f"The PDF has {len(pdf)} pages; the limit is {max_pages} (MAX_PDF_PAGES)")
        pages = []
        for i in range(len(pdf)):
            page = pdf[i]
            try:
                text_page = page.get_textpage()
                try:
                    pages.append(text_page.get_text_range())
                finally:
                    text_page.close()
            finally:
                page.close()
        return pages
    finally:
        pdf.close()


def ocr_pdf_pages(path: Path, indexes: list[int]) -> dict[int, str]:
    """OCR the given 0-based pages (render with pdfium, tesseract with the Indonesian model)."""
    import pytesseract

    try:
        pytesseract.get_tesseract_version()
    except pytesseract.TesseractNotFoundError as e:
        raise PermanentError(f"{len(indexes)} page(s) need OCR but tesseract is not installed") from e
    pdf = pdfium.PdfDocument(str(path))
    out: dict[int, str] = {}
    try:
        for done, i in enumerate(indexes):
            if done and done % 25 == 0:
                log.info("OCR %s: %d/%d pages", path.name, done, len(indexes))
            page = pdf[i]
            try:
                scale = min(OCR_DPI / 72, OCR_MAX_PIXELS / max(page.get_size()))
                image = page.render(scale=scale).to_pil()
            finally:
                page.close()
            try:
                out[i] = pytesseract.image_to_string(image, lang="ind", timeout=OCR_PAGE_TIMEOUT)
            except pytesseract.TesseractError as e:
                log.error("tesseract failed on %s: %s", path.name, e)
                raise PermanentError("OCR failed (is the tesseract 'ind' language data installed?)") from e
            except RuntimeError:  # pytesseract signals a per-page timeout this way: skip the page
                log.warning("OCR of page %d of %s timed out; skipped", i + 1, path.name)
        return out
    finally:
        pdf.close()


def _check_docx_zip(path: Path) -> None:
    """A DOCX is a zip that python-docx unpacks fully into memory: refuse a zip bomb from the archive index, before
    opening it. (zipfile never inflates past the size the index declares, so the index figures are enforceable.)"""
    try:
        with zipfile.ZipFile(path) as z:
            infos = z.infolist()
    except (zipfile.BadZipFile, OSError) as e:
        log.warning("cannot read %s as a zip: %s", path.name, e)
        raise PermanentError("Cannot read DOCX (not a valid or readable Word file)") from e
    total, packed = sum(i.file_size for i in infos), sum(i.compress_size for i in infos)
    if total > DOCX_MAX_UNCOMPRESSED:
        raise PermanentError(f"The DOCX unpacks to {total >> 20} MB; the limit is {DOCX_MAX_UNCOMPRESSED >> 20} MB")
    if total > DOCX_RATIO_FLOOR and total > packed * DOCX_MAX_RATIO:
        raise PermanentError(f"The DOCX is compressed more than {DOCX_MAX_RATIO}:1, which looks like a zip bomb")


def _read_docx(path: Path) -> str:
    import docx
    from docx.table import Table

    _check_docx_zip(path)
    try:
        document = docx.Document(str(path))
        lines = []
        for block in document.iter_inner_content():
            if isinstance(block, Table):
                lines += [" | ".join(cell.text.strip() for cell in row.cells) for row in block.rows]
            else:
                heading = block.style is not None and block.style.name.startswith("Heading")
                lines.append(f"# {block.text}" if heading and block.text.strip() else block.text)
    except Exception as e:  # python-docx raises many unrelated types on a bad file
        log.warning("cannot read %s as DOCX: %s: %s", path.name, type(e).__name__, e)
        raise PermanentError("Cannot read DOCX (not a valid or readable Word file)") from e
    return "\n".join(lines)


def read_document(path: Path, ext: str, settings: Settings) -> Extracted:
    """Extract and clean the text. For a PDF, pages that stay under OCR_MIN_CHARS after cleaning are OCRed
    (the OCR text is cleaned the same way and kept only when it is longer than the text layer). A PDF over
    MAX_PDF_PAGES pages, or needing OCR on more than MAX_OCR_PAGES, fails instead of tying the worker up."""
    if ext == "pdf":
        pages = [clean_page(t) for t in extract_pdf(path, settings.max_pdf_pages)]
        todo = [i for i, t in enumerate(pages) if len(t) < OCR_MIN_CHARS]
        if len(todo) > settings.max_ocr_pages:
            raise PermanentError(f"{len(todo)} pages need OCR; the limit is {settings.max_ocr_pages} (MAX_OCR_PAGES)")
        used = 0
        if todo:
            for i, text in ocr_pdf_pages(path, todo).items():
                text = clean_page(text)
                if len(text) > len(pages[i]):
                    pages[i], used = text, used + 1
        return Extracted(pages, True, used)
    text = _read_docx(path) if ext == "docx" else path.read_text(encoding="utf-8", errors="replace")
    return Extracted(["\n".join(normalize(text))], False)


# --- job queue -------------------------------------------------------------------------------------------


def claim_job(conn: Connection) -> dict | None:
    """Atomically take the oldest queued job (FOR UPDATE SKIP LOCKED) and count the attempt."""
    return conn.execute(
        "UPDATE ingestion_jobs SET status = 'running', attempts = attempts + 1, started_at = now(), error = NULL "
        "WHERE id = (SELECT id FROM ingestion_jobs WHERE status = 'queued' ORDER BY id "
        "FOR UPDATE SKIP LOCKED LIMIT 1) RETURNING id, kb_id, document_id, attempts"
    ).fetchone()


def reset_stale(conn: Connection) -> int:
    """At worker start: jobs left 'running' by a dead worker go back to the queue (or fail once out of attempts).

    ponytail: assumes a single worker. With several, a lease timeout (heartbeat column) replaces this.
    """
    rows = conn.execute(
        "WITH j AS (UPDATE ingestion_jobs SET status = CASE WHEN attempts >= %(max)s THEN 'failed' ELSE 'queued' END, "
        "error = CASE WHEN attempts >= %(max)s THEN 'Worker stopped while processing' ELSE error END "
        "WHERE status = 'running' RETURNING document_id, kb_id, status, error) "
        "UPDATE documents d SET status = j.status, error = j.error FROM j "
        "WHERE d.id = j.document_id AND d.kb_id = j.kb_id RETURNING d.id",
        {"max": MAX_ATTEMPTS},
    ).fetchall()
    return len(rows)


# stats of a document's latest finished job; `d` is the documents alias of the enclosing query
_LATEST_STATS = (
    "(SELECT j.stats FROM ingestion_jobs j WHERE j.document_id = d.id AND j.kb_id = d.kb_id AND j.status = 'done' "
    "ORDER BY j.id DESC LIMIT 1)"
)


def maybe_switch_model(conn: Connection, kb_id: UUID, settings: Settings) -> bool:
    """Make the configured embedding setup the KB's once every document that holds vectors was last indexed with it.

    Until then a KB indexed with another setup keeps answering "reindex required", so it never serves a mix of
    two vector spaces. The fingerprint of each finished job is kept in its stats (hidden by the jobs API). A
    document without chunks (failed before it ever indexed) holds no vectors and does not hold the switch back.
    """
    return (
        conn.execute(
            "UPDATE knowledge_bases kb SET embedding_model = %(m)s, embedding_fingerprint = %(fp)s, updated_at = now() "
            "WHERE kb.id = %(kb)s AND kb.embedding_fingerprint IS DISTINCT FROM %(fp)s "
            "AND kb.embedding_canary IS NOT NULL AND EXISTS (SELECT 1 FROM documents WHERE kb_id = kb.id) "
            "AND NOT EXISTS (SELECT 1 FROM documents d WHERE d.kb_id = kb.id AND d.chunk_count > 0 "
            f"AND {_LATEST_STATS}->>'embedding_fingerprint' IS DISTINCT FROM %(fp)s)",
            {"m": settings.embedding_model, "fp": settings.embedding_fingerprint, "kb": kb_id},
        ).rowcount
        > 0
    )


def stray_setup(conn: Connection, kb_id: UUID, fingerprint: str) -> dict | None:
    """(model, fingerprint) of some document that holds vectors but was last indexed with another setup than
    `fingerprint`, or None. A finished job without a recorded fingerprint counts as 'unknown'."""
    return conn.execute(
        f"SELECT coalesce({_LATEST_STATS}->>'embedding_model', 'unknown') AS model, "
        f"coalesce({_LATEST_STATS}->>'embedding_fingerprint', 'unknown') AS fingerprint "
        "FROM documents d WHERE d.kb_id = %(kb)s AND d.chunk_count > 0 "
        f"AND {_LATEST_STATS}->>'embedding_fingerprint' IS DISTINCT FROM %(fp)s LIMIT 1",
        {"kb": kb_id, "fp": fingerprint},
    ).fetchone()


def _finish_failed(pool: ConnectionPool, job: dict, message: str, retry: bool) -> None:
    with pool.connection() as conn, conn.transaction():
        conn.execute(
            "UPDATE ingestion_jobs SET status = %s, error = %s, finished_at = now() WHERE id = %s",
            ["queued" if retry else "failed", message, job["id"]],
        )
        conn.execute(
            "UPDATE documents SET status = %s, error = %s WHERE id = %s AND kb_id = %s",
            ["queued" if retry else "failed", message, job["document_id"], job["kb_id"]],
        )


def _embed_text(title: str, chunk: Chunk) -> str:
    """The text that is embedded: title and heading in front of the content (the stored content stays clean)."""
    return "\n".join(part for part in (title, chunk.heading, chunk.content) if part)


def _public_message(e: Exception) -> str:
    """What goes into documents.error and ingestion_jobs.error, which every KB member can read. Only messages our own
    code wrote pass (PermanentError, and EmbeddingError, whose text is sanitized at the source); the detail of
    anything else stays in the server log."""
    if isinstance(e, PermanentError | embedding.EmbeddingError):
        return str(e)[:1000]
    return f"Unexpected error ({type(e).__name__}); see the server log"


def process_job(pool: ConnectionPool, job: dict, settings: Settings) -> None:
    """Run one claimed job to completion. Never raises: the outcome is recorded on the job and document rows."""
    started = time.monotonic()
    try:
        with pool.connection() as conn:
            doc = conn.execute(
                "SELECT d.filename, d.storage_path, kb.embedding_model, kb.embedding_fingerprint, "
                "kb.embedding_canary, kb.embedding_query_canary "
                "FROM documents d JOIN knowledge_bases kb ON kb.id = d.kb_id WHERE d.id = %s AND d.kb_id = %s",
                [job["document_id"], job["kb_id"]],
            ).fetchone()
            if doc is None:
                log.info("job %s: document was deleted, nothing to do", job["id"])
                return
            conn.execute(
                "UPDATE documents SET status = 'processing', error = NULL WHERE id = %s AND kb_id = %s",
                [job["document_id"], job["kb_id"]],
            )
        # No connection is held below: extraction, OCR and embedding can take minutes.
        path = resolve_stored(settings, doc["storage_path"])
        ext = path.suffix.lstrip(".").lower()
        if ext not in ALLOWED_TYPES:
            raise PermanentError(f"Unsupported file type: .{ext}")
        if not path.is_file():
            raise PermanentError("The stored file is missing")
        extracted = read_document(path, ext, settings)
        chunks = chunk_pages(extracted.pages, extracted.paged)
        if not chunks:
            raise PermanentError("No text could be extracted from this document")
        title = Path(doc["filename"]).stem
        try:
            result = embedding.embed(
                [_embed_text(title, c) for c in chunks], "document", canary=doc["embedding_canary"], settings=settings
            )
        except embedding.EmbeddingMismatch as e:
            kb_fingerprint = doc["embedding_fingerprint"]
            if kb_fingerprint and kb_fingerprint != settings.embedding_fingerprint:
                raise PermanentError(
                    f"The embedding setup (model '{settings.embedding_model}', prefixes or extra body) differs from "
                    f"the one this knowledge base was indexed with ('{doc['embedding_model']}'); reindex the whole "
                    f"knowledge base first ({e})"
                ) from e
            raise
        # The query-mode canary is captured once, next to the document one (first index, or after a reset).
        query_canary = (
            embedding.query_canary(result.canary, settings) if doc["embedding_query_canary"] is None else None
        )
        stats = {
            "chunks": len(chunks),
            "ocr_pages": extracted.ocr_pages,
            "tokens": result.tokens,
            "seconds": round(time.monotonic() - started, 1),
            "embedding_model": settings.embedding_model,
            "embedding_fingerprint": settings.embedding_fingerprint,
        }
        if extracted.paged:
            stats["pages"] = len(extracted.pages)
        rows = [
            (job["kb_id"], job["document_id"], c.index, c.page_start, c.page_end, c.heading, c.content, Vector(v))
            for c, v in zip(chunks, result.vectors, strict=True)
        ]
        with pool.connection() as conn:
            with kb_scope(conn, job["kb_id"]):  # one transaction: chunks, document, job and KB canary together
                conn.execute(
                    "DELETE FROM document_chunks WHERE document_id = %s AND kb_id = %s",
                    [job["document_id"], job["kb_id"]],
                )
                conn.cursor().executemany(
                    "INSERT INTO document_chunks (kb_id, document_id, chunk_index, page_start, page_end, heading, "
                    "content, embedding) VALUES (%s, %s, %s, %s, %s, %s, %s, %s)",
                    rows,
                )
                conn.execute(
                    "UPDATE documents SET status = 'ready', page_count = %s, ocr_pages = %s, chunk_count = %s, "
                    "error = NULL, indexed_at = now() WHERE id = %s AND kb_id = %s",
                    [
                        len(extracted.pages) if extracted.paged else None,
                        extracted.ocr_pages,
                        len(chunks),
                        job["document_id"],
                        job["kb_id"],
                    ],
                )
                conn.execute(
                    "UPDATE ingestion_jobs SET status = 'done', stats = %s, error = NULL, finished_at = now() "
                    "WHERE id = %s",
                    [Jsonb(stats), job["id"]],
                )
                # First index sets the setup and both canaries; during a model change only the (reset) canaries.
                conn.execute(
                    "UPDATE knowledge_bases SET embedding_model = coalesce(embedding_model, %s), "
                    "embedding_fingerprint = coalesce(embedding_fingerprint, %s), "
                    "embedding_canary = coalesce(embedding_canary, %s), "
                    "embedding_query_canary = coalesce(embedding_query_canary, %s::vector) WHERE id = %s",
                    [
                        settings.embedding_model,
                        settings.embedding_fingerprint,
                        Vector(result.canary),
                        None if query_canary is None else Vector(query_canary),
                        job["kb_id"],
                    ],
                )
            maybe_switch_model(conn, job["kb_id"], settings)
        log.info(
            "job %s done: %d chunks, %d OCR pages, %.1fs",
            job["id"],
            len(chunks),
            extracted.ocr_pages,
            time.monotonic() - started,
        )
    except Exception as e:  # every failure must land on the job row, not kill the worker
        permanent = getattr(e, "permanent", False)
        retry = not permanent and job["attempts"] < MAX_ATTEMPTS
        log.exception("job %s failed (attempt %d, %s)", job["id"], job["attempts"], "retrying" if retry else "final")
        try:
            _finish_failed(pool, job, _public_message(e), retry)
        except Exception:
            log.exception("job %s: could not record the failure", job["id"])
